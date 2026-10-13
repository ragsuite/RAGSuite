"""Heading outlines of uploaded documents (PDF bookmarks, DOCX heading styles, PPTX titles).

Every reader is best-effort: a document without an outline (or one that fails to
parse) simply yields no headings, and chunks fall back to title and breadcrumb.
"""
from __future__ import annotations

import logging
from typing import Any, Dict, List

logger = logging.getLogger(__name__)

MAX_OUTLINE_ITEMS = 200
MAX_HEADING_CHARS = 120


def _clean(value: Any) -> str:
    text = " ".join(str(value or "").split())
    return text if 2 <= len(text) <= MAX_HEADING_CHARS else ""


def pdf_outline_titles(reader: Any) -> List[str]:
    """Bookmark titles of a ``pypdf.PdfReader`` in reading order."""
    titles: List[str] = []

    def walk(items: Any) -> None:
        for item in items or []:
            if len(titles) >= MAX_OUTLINE_ITEMS:
                return
            if isinstance(item, list):
                walk(item)
                continue
            title = _clean(getattr(item, "title", None))
            if title:
                titles.append(title)

    try:
        walk(reader.outline)
    except Exception as exc:
        logger.debug("pdf outline unavailable: %s", exc)
    return titles


def docx_heading_texts(filepath: str) -> List[str]:
    """Paragraphs styled ``Title`` or ``Heading n`` in a .docx file."""
    try:
        import docx

        document = docx.Document(filepath)
    except Exception as exc:
        logger.debug("docx outline unavailable for %s: %s", filepath, exc)
        return []
    headings: List[str] = []
    for paragraph in document.paragraphs:
        style = (getattr(paragraph.style, "name", "") or "").lower()
        if style == "title" or style.startswith("heading"):
            text = _clean(paragraph.text)
            if text:
                headings.append(text)
        if len(headings) >= MAX_OUTLINE_ITEMS:
            break
    return headings


def pptx_slide_titles(filepath: str) -> Dict[int, str]:
    """Slide number (1-based) to title placeholder text."""
    try:
        from pptx import Presentation

        presentation = Presentation(filepath)
    except Exception as exc:
        logger.debug("pptx titles unavailable for %s: %s", filepath, exc)
        return {}
    titles: Dict[int, str] = {}
    for number, slide in enumerate(presentation.slides, start=1):
        try:
            shape = slide.shapes.title
        except Exception:
            shape = None
        title = _clean(getattr(shape, "text", "") if shape is not None else "")
        if title:
            titles[number] = title
    return titles
