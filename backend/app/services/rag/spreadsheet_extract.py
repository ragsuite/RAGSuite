"""
XLSX extraction for RAG ingest.

Each visible sheet becomes header-labelled row lines (``Name: Ada; Role: CTO``)
packed into sheet-scoped chunks. Rows are never split across chunks unless a
single row exceeds the chunk size on its own.
"""
from __future__ import annotations

import datetime as dt
import logging
import os
import re
from dataclasses import dataclass, field
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple

logger = logging.getLogger(__name__)

MAX_XLSX_DATA_ROWS = 50_000
_MIN_PART_SIZE = 400
_WHITESPACE_RE = re.compile(r"\s+")

SplitText = Callable[[str, int], List[str]]


def _clean(text: str) -> str:
    collapsed = _WHITESPACE_RE.sub(" ", text).strip()
    # Lone surrogates break Chroma/JSON serialization.
    return collapsed.encode("utf-8", "ignore").decode("utf-8")


def format_cell(value: Any) -> str:
    """Render a cell value as compact, embedding-friendly text."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds" if value.second else "minutes")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return _clean(str(value))


def _column_letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index + 1)


def _is_header_row(values: Sequence[str], raw: Sequence[Any]) -> bool:
    filled = [r for r, v in zip(raw, values) if v]
    return bool(filled) and all(isinstance(r, str) for r in filled)


def _row_text(values: Sequence[str], labels: Sequence[str]) -> str:
    parts = []
    for idx, val in enumerate(values):
        if not val:
            continue
        label = labels[idx] if idx < len(labels) and labels[idx] else _column_letter(idx)
        parts.append(f"{label}: {val}")
    return "; ".join(parts)


@dataclass
class _ChunkBuilder:
    title: str
    url: str
    chunk_size: int
    split_text: SplitText
    texts: List[str] = field(default_factory=list)
    metadata: List[Dict[str, Any]] = field(default_factory=list)

    def emit(self, text: str, sheet_name: str, sheet_index: int, row_start: int, row_end: int) -> None:
        self.texts.append(text)
        self.metadata.append(
            {
                "title": self.title,
                "url": self.url,
                "keywords": "",
                "chunk_index": len(self.texts) - 1,
                "sheet_name": sheet_name,
                "sheet_index": sheet_index,
                "page": sheet_index,
                "row_start": row_start,
                "row_end": row_end,
                "truncated": False,
            }
        )

    def add_sheet(self, sheet_name: str, sheet_index: int, rows: List[Tuple[int, str]]) -> None:
        prefix = f"Sheet: {sheet_name}"
        lines: List[str] = []
        first_row = last_row = 0
        size = len(prefix)

        def flush() -> None:
            nonlocal lines, size
            if lines:
                self.emit("\n".join([prefix, *lines]), sheet_name, sheet_index, first_row, last_row)
            lines, size = [], len(prefix)

        for row_num, text in rows:
            line = f"Row {row_num}: {text}"
            if len(prefix) + 1 + len(line) > self.chunk_size:
                flush()
                self._emit_oversized_row(prefix, sheet_name, sheet_index, row_num, text)
                continue
            if lines and size + 1 + len(line) > self.chunk_size:
                flush()
            if not lines:
                first_row = row_num
            lines.append(line)
            last_row = row_num
            size += 1 + len(line)
        flush()

    def _emit_oversized_row(
        self, prefix: str, sheet_name: str, sheet_index: int, row_num: int, text: str
    ) -> None:
        part_budget = max(_MIN_PART_SIZE, self.chunk_size - len(prefix) - 32)
        parts = [p for p in self.split_text(text, part_budget) if p.strip()] or [text]
        total = len(parts)
        for idx, part in enumerate(parts, start=1):
            label = f"Row {row_num}" if total == 1 else f"Row {row_num} (part {idx}/{total})"
            self.emit(f"{prefix}\n{label}: {part}", sheet_name, sheet_index, row_num, row_num)


def _load_workbook(filepath: str):
    try:
        from openpyxl import load_workbook

        return load_workbook(filepath, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Not a valid XLSX workbook") from exc


def _sheet_rows(ws, budget: int) -> Tuple[List[Tuple[int, str]], bool]:
    """Return ``[(excel_row_number, row_text), ...]`` and whether the budget ran out."""
    ws.reset_dimensions()
    labels: Optional[List[str]] = None
    rows: List[Tuple[int, str]] = []
    for row_num, raw in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1):
        raw = raw or ()
        values = [format_cell(v) for v in raw]
        if not any(values):
            continue
        if labels is None:
            if _is_header_row(values, raw):
                labels = values
                continue
            labels = []
        if len(rows) >= budget:
            return rows, True
        text = _row_text(values, labels)
        if text:
            rows.append((row_num, text))
    return rows, False


def extract_xlsx_chunks(
    filepath: str,
    *,
    chunk_size: int,
    split_text: SplitText,
    max_rows: Optional[int] = None,
) -> Tuple[List[str], List[Dict[str, Any]]]:
    """Return ``(texts, metadata)`` for an ``.xlsx`` workbook (visible sheets only)."""
    row_cap = MAX_XLSX_DATA_ROWS if max_rows is None else max_rows
    name = os.path.basename(filepath)
    builder = _ChunkBuilder(
        title=os.path.splitext(name)[0],
        url=f"file://{name}",
        chunk_size=chunk_size,
        split_text=split_text,
    )
    wb = _load_workbook(filepath)
    truncated = False
    rows_used = 0
    try:
        for sheet_index, ws in enumerate(wb.worksheets, start=1):
            if getattr(ws, "sheet_state", "visible") != "visible":
                continue
            rows, truncated = _sheet_rows(ws, row_cap - rows_used)
            rows_used += len(rows)
            builder.add_sheet(_clean(ws.title) or f"Sheet{sheet_index}", sheet_index, rows)
            if truncated:
                break
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Not a valid XLSX workbook") from exc
    finally:
        wb.close()

    if truncated:
        logger.warning(
            "XLSX %s truncated at %s data rows; remaining rows were not indexed", name, row_cap
        )
        for meta in builder.metadata:
            meta["truncated"] = True
    return builder.texts, builder.metadata
