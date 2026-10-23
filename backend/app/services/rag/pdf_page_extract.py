"""Page/slide/sheet range and header/footer strip for uploaded documents.

Extract defaults stay all-units / no strip. Uploaded-document ingest passes options
from ``UploadedDocument.meta_data``. Crawl and connectors do not use this module.

Meta keys stay ``pdf_*`` so already-uploaded PDFs keep working. The same keys
apply to PPTX (slides), XLSX (sheets), and DOCX (Word header/footer strip).
"""
from __future__ import annotations

import math
import os
import re
from dataclasses import dataclass
from typing import Any, Dict, List, Optional, Sequence, Tuple

from .docling_extract import (
    ENGINE_DOCLING,
    ENGINE_NATIVE,
    EXTRACT_ENGINE_KEY,
    filename_supports_docling,
    normalize_extract_engine,
)

RANGE_EXTS = frozenset({".pdf", ".pptx", ".xlsx"})
STRIP_EXTS = frozenset({".pdf", ".pptx", ".docx"})

PDF_PAGE_FROM_KEY = "pdf_page_from"
PDF_PAGE_TO_KEY = "pdf_page_to"
PDF_STRIP_CHROME_KEY = "pdf_strip_chrome"
PDF_STRIP_HEADER_KEY = "pdf_strip_header"
PDF_STRIP_FOOTER_KEY = "pdf_strip_footer"
OCR_MODEL_KEY = "ocr_model"

UNSET = object()

_HEADER_FOOTER_LINES = 3
_PAGE_NUMBER_RE = re.compile(
    r"^(?:page\s+\d+(?:\s*(?:of|/)\s*\d+)?|\d+\s*/\s*\d+|[-–—]\s*\d+\s*[-–—])$",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class PdfExtractOptions:
    page_from: Optional[int] = None
    page_to: Optional[int] = None
    strip_header: bool = False
    strip_footer: bool = False

    @property
    def strip_chrome(self) -> bool:
        return self.strip_header or self.strip_footer


def _optional_bool(value: Any) -> Optional[bool]:
    if value is None:
        return None
    return bool(value)


def resolve_strip_sides(
    *,
    strip_chrome: Any = None,
    strip_header: Any = None,
    strip_footer: Any = None,
    default_when_unset: bool = False,
) -> Tuple[bool, bool]:
    """Independent header/footer. Legacy chrome alone applies to both sides."""
    header = _optional_bool(strip_header)
    footer = _optional_bool(strip_footer)
    chrome = _optional_bool(strip_chrome)
    if header is None and footer is None:
        on = default_when_unset if chrome is None else chrome
        return on, on
    fallback = default_when_unset if chrome is None else chrome
    return (
        fallback if header is None else header,
        fallback if footer is None else footer,
    )


def _leaf_name(name: Optional[str]) -> str:
    return (name or "").split("?")[0].strip().lower()


def filename_ext(name: Optional[str]) -> str:
    return os.path.splitext(_leaf_name(name))[1]


def is_pdf_filename(name: Optional[str]) -> bool:
    return filename_ext(name) == ".pdf"


def is_zip_filename(name: Optional[str]) -> bool:
    return filename_ext(name) == ".zip"


def filename_supports_range(name: Optional[str]) -> bool:
    ext = filename_ext(name)
    return ext in RANGE_EXTS or ext == ".zip"


def filename_supports_strip(name: Optional[str]) -> bool:
    ext = filename_ext(name)
    return ext in STRIP_EXTS or ext == ".zip"


def document_ext(doc: Any) -> str:
    for candidate in (getattr(doc, "title", None), getattr(doc, "source", None)):
        ext = filename_ext(candidate)
        if ext in RANGE_EXTS | STRIP_EXTS | {".doc", ".zip"}:
            return ext
    mime = (getattr(doc, "type", None) or "").lower()
    if "pdf" in mime:
        return ".pdf"
    if "presentation" in mime or "pptx" in mime:
        return ".pptx"
    if "spreadsheet" in mime or "xlsx" in mime or "excel" in mime:
        return ".xlsx"
    if "wordprocessingml" in mime or "docx" in mime:
        return ".docx"
    if "msword" in mime or mime.endswith("/doc"):
        return ".doc"
    return ""


def is_pdf_document(doc: Any) -> bool:
    return document_ext(doc) == ".pdf"


def supports_extract_options(doc: Any) -> bool:
    ext = document_ext(doc)
    return ext in RANGE_EXTS or ext in STRIP_EXTS or ext == ".zip"


def validate_pdf_page_range(
    page_from: Any, page_to: Any
) -> Tuple[Optional[int], Optional[int]]:
    """1-indexed inclusive range. Empty/None = all pages. Raises ValueError."""

    def _parse(value: Any, name: str) -> Optional[int]:
        if value is None or value == "":
            return None
        try:
            number = int(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"{name} must be a positive integer") from exc
        if number < 1:
            raise ValueError(f"{name} must be >= 1")
        return number

    start = _parse(page_from, PDF_PAGE_FROM_KEY)
    end = _parse(page_to, PDF_PAGE_TO_KEY)
    if start is not None and end is not None and start > end:
        raise ValueError("pdf_page_from must be less than or equal to pdf_page_to")
    return start, end


def options_from_meta(meta: Optional[Dict[str, Any]]) -> PdfExtractOptions:
    """Missing keys = current extract (all pages, no chrome strip)."""
    data = meta if isinstance(meta, dict) else {}
    try:
        start, end = validate_pdf_page_range(
            data.get(PDF_PAGE_FROM_KEY), data.get(PDF_PAGE_TO_KEY)
        )
    except ValueError:
        start, end = None, None
    header, footer = resolve_strip_sides(
        strip_chrome=data.get(PDF_STRIP_CHROME_KEY),
        strip_header=data.get(PDF_STRIP_HEADER_KEY) if PDF_STRIP_HEADER_KEY in data else None,
        strip_footer=data.get(PDF_STRIP_FOOTER_KEY) if PDF_STRIP_FOOTER_KEY in data else None,
    )
    return PdfExtractOptions(
        page_from=start,
        page_to=end,
        strip_header=header,
        strip_footer=footer,
    )


def ingest_kwargs_from_options(opts: PdfExtractOptions) -> Dict[str, Any]:
    kwargs: Dict[str, Any] = {}
    if opts.page_from is not None:
        kwargs["pdf_page_from"] = opts.page_from
    if opts.page_to is not None:
        kwargs["pdf_page_to"] = opts.page_to
    if opts.strip_header or opts.strip_footer:
        kwargs["pdf_strip_chrome"] = True
        kwargs["pdf_strip_header"] = opts.strip_header
        kwargs["pdf_strip_footer"] = opts.strip_footer
    return kwargs


def extract_engine_from_meta(meta: Optional[Dict[str, Any]]) -> str:
    data = meta if isinstance(meta, dict) else {}
    return normalize_extract_engine(data.get(EXTRACT_ENGINE_KEY))


def ocr_model_from_meta(meta: Optional[Dict[str, Any]]) -> Optional[str]:
    data = meta if isinstance(meta, dict) else {}
    value = str(data.get(OCR_MODEL_KEY) or "").strip()
    return value or None


def ingest_kwargs_from_document(doc: Any) -> Dict[str, Any]:
    meta = getattr(doc, "meta_data", None)
    kwargs = ingest_kwargs_from_options(options_from_meta(meta))
    engine = extract_engine_from_meta(meta)
    if engine != ENGINE_NATIVE:
        kwargs["extract_engine"] = engine
    ocr_model = ocr_model_from_meta(meta)
    if ocr_model:
        kwargs["ocr_model"] = ocr_model
        kwargs["extract_engine"] = ENGINE_DOCLING if engine == ENGINE_NATIVE else engine
    from .split_config import snapshot_split_config

    stored = snapshot_split_config(meta)
    if stored:
        kwargs["split_config"] = stored
    return kwargs


def pdf_fields_for_api(meta: Optional[Dict[str, Any]]) -> Dict[str, Optional[Any]]:
    data = meta if isinstance(meta, dict) else {}
    opts = options_from_meta(data)
    header_raw = data.get(PDF_STRIP_HEADER_KEY) if PDF_STRIP_HEADER_KEY in data else None
    footer_raw = data.get(PDF_STRIP_FOOTER_KEY) if PDF_STRIP_FOOTER_KEY in data else None
    strip_raw = data.get(PDF_STRIP_CHROME_KEY) if PDF_STRIP_CHROME_KEY in data else None
    header_out = None if header_raw is None and strip_raw is None else opts.strip_header
    footer_out = None if footer_raw is None and strip_raw is None else opts.strip_footer
    chrome_out = None if header_out is None and footer_out is None and strip_raw is None else opts.strip_chrome
    return {
        "pdfPageFrom": opts.page_from,
        "pdfPageTo": opts.page_to,
        "pdfStripChrome": chrome_out,
        "pdfStripHeader": header_out,
        "pdfStripFooter": footer_out,
        "extractEngine": extract_engine_from_meta(data) if EXTRACT_ENGINE_KEY in data else None,
        "ocrModel": ocr_model_from_meta(data),
    }


def meta_for_upload(
    *,
    page_from: Any,
    page_to: Any,
    strip_chrome: Any,
    filename: Optional[str],
    extract_engine: Any = None,
    strip_header: Any = None,
    strip_footer: Any = None,
    ocr_model: Any = None,
) -> Dict[str, Any]:
    store_range = filename_supports_range(filename)
    store_strip = filename_supports_strip(filename)
    store_engine = filename_supports_docling(filename)
    if not store_range and not store_strip and not store_engine:
        return {}
    start, end = validate_pdf_page_range(page_from, page_to)
    meta: Dict[str, Any] = {}
    if store_strip:
        header, footer = resolve_strip_sides(
            strip_chrome=strip_chrome,
            strip_header=strip_header,
            strip_footer=strip_footer,
            default_when_unset=True,
        )
        meta[PDF_STRIP_HEADER_KEY] = header
        meta[PDF_STRIP_FOOTER_KEY] = footer
        meta[PDF_STRIP_CHROME_KEY] = header or footer
    if store_range:
        if start is not None:
            meta[PDF_PAGE_FROM_KEY] = start
        if end is not None:
            meta[PDF_PAGE_TO_KEY] = end
    if store_engine:
        meta[EXTRACT_ENGINE_KEY] = normalize_extract_engine(extract_engine)
        value = str(ocr_model or "").strip()
        if value:
            meta[OCR_MODEL_KEY] = value
            meta[EXTRACT_ENGINE_KEY] = ENGINE_DOCLING
    return meta


def apply_pdf_options_to_meta(
    meta: Optional[Dict[str, Any]],
    *,
    page_from: Any = UNSET,
    page_to: Any = UNSET,
    strip_chrome: Any = UNSET,
    extract_engine: Any = UNSET,
    strip_header: Any = UNSET,
    strip_footer: Any = UNSET,
    ocr_model: Any = UNSET,
) -> Tuple[Dict[str, Any], bool]:
    current = dict(meta) if isinstance(meta, dict) else {}
    before = options_from_meta(current)
    before_engine = extract_engine_from_meta(current)
    if page_from is not UNSET or page_to is not UNSET:
        new_from = before.page_from if page_from is UNSET else page_from
        new_to = before.page_to if page_to is UNSET else page_to
        start, end = validate_pdf_page_range(new_from, new_to)
        if start is None:
            current.pop(PDF_PAGE_FROM_KEY, None)
        else:
            current[PDF_PAGE_FROM_KEY] = start
        if end is None:
            current.pop(PDF_PAGE_TO_KEY, None)
        else:
            current[PDF_PAGE_TO_KEY] = end
    if strip_chrome is not UNSET:
        if strip_chrome is None:
            current.pop(PDF_STRIP_CHROME_KEY, None)
        else:
            current[PDF_STRIP_CHROME_KEY] = bool(strip_chrome)
            if strip_header is UNSET:
                current[PDF_STRIP_HEADER_KEY] = bool(strip_chrome)
            if strip_footer is UNSET:
                current[PDF_STRIP_FOOTER_KEY] = bool(strip_chrome)
    if strip_header is not UNSET:
        if strip_header is None:
            current.pop(PDF_STRIP_HEADER_KEY, None)
        else:
            current[PDF_STRIP_HEADER_KEY] = bool(strip_header)
    if strip_footer is not UNSET:
        if strip_footer is None:
            current.pop(PDF_STRIP_FOOTER_KEY, None)
        else:
            current[PDF_STRIP_FOOTER_KEY] = bool(strip_footer)
    if PDF_STRIP_HEADER_KEY in current or PDF_STRIP_FOOTER_KEY in current:
        resolved = options_from_meta(current)
        current[PDF_STRIP_CHROME_KEY] = resolved.strip_header or resolved.strip_footer
    if extract_engine is not UNSET:
        if extract_engine is None or extract_engine == "":
            current.pop(EXTRACT_ENGINE_KEY, None)
        else:
            current[EXTRACT_ENGINE_KEY] = normalize_extract_engine(extract_engine)
    before_ocr = ocr_model_from_meta(current)
    if ocr_model is not UNSET:
        value = str(ocr_model or "").strip()
        if value:
            current[OCR_MODEL_KEY] = value
            current[EXTRACT_ENGINE_KEY] = ENGINE_DOCLING
        else:
            current.pop(OCR_MODEL_KEY, None)
    after = options_from_meta(current)
    after_engine = extract_engine_from_meta(current)
    after_ocr = ocr_model_from_meta(current)
    changed = (
        before.page_from,
        before.page_to,
        before.strip_header,
        before.strip_footer,
        before_engine,
        before_ocr,
    ) != (
        after.page_from,
        after.page_to,
        after.strip_header,
        after.strip_footer,
        after_engine,
        after_ocr,
    )
    return current, changed


def _nonempty_lines(text: str) -> List[str]:
    return [line.strip() for line in (text or "").splitlines() if line.strip()]


def _norm_line(line: str) -> str:
    return " ".join((line or "").split()).lower()


def _is_page_number_line(line: str) -> bool:
    return bool(_PAGE_NUMBER_RE.match(_norm_line(line)))


def _edge_candidate_lines(
    lines: Sequence[str], *, strip_header: bool, strip_footer: bool
) -> List[str]:
    values = list(lines)
    if strip_header and strip_footer:
        edge = list(values[:_HEADER_FOOTER_LINES])
        if len(values) > _HEADER_FOOTER_LINES:
            edge.extend(values[-_HEADER_FOOTER_LINES:])
        elif values:
            edge = list(values)
        return edge
    edge: List[str] = []
    if strip_header:
        edge.extend(values[:_HEADER_FOOTER_LINES])
    if strip_footer:
        if len(values) > _HEADER_FOOTER_LINES:
            edge.extend(values[-_HEADER_FOOTER_LINES:])
        elif values:
            edge.extend(values)
    return edge


def strip_repeating_chrome(
    page_texts: Sequence[str],
    *,
    strip_header: bool = True,
    strip_footer: bool = True,
) -> List[str]:
    if not strip_header and not strip_footer:
        return [text or "" for text in page_texts]
    parsed = [_nonempty_lines(text) for text in page_texts]
    page_count = len(parsed)
    counts: Dict[str, int] = {}
    for lines in parsed:
        seen = set()
        for line in _edge_candidate_lines(
            lines, strip_header=strip_header, strip_footer=strip_footer
        ):
            key = _norm_line(line)
            if not key or key in seen:
                continue
            seen.add(key)
            counts[key] = counts.get(key, 0) + 1
    threshold = max(3, math.ceil(page_count * 0.5)) if page_count else 3
    chrome = {key for key, count in counts.items() if count >= threshold}

    cleaned: List[str] = []
    for lines in parsed:
        kept = [line for line in lines if _norm_line(line) not in chrome]
        if strip_header:
            while kept and _is_page_number_line(kept[0]):
                kept.pop(0)
        if strip_footer:
            while kept and _is_page_number_line(kept[-1]):
                kept.pop()
        cleaned.append("\n".join(kept))
    return cleaned


def clamp_unit_range(
    total: int, page_from: Optional[int], page_to: Optional[int]
) -> Tuple[int, int]:
    """1-indexed inclusive range clamped to ``total`` units."""
    if total <= 0:
        return 1, 0
    start = 1 if page_from is None else int(page_from)
    end = total if page_to is None else int(page_to)
    start = max(1, min(start, total))
    end = max(1, min(end, total))
    if start > end:
        start, end = end, start
    return start, end


def slice_numbered_texts(
    items: Sequence[Tuple[int, str]],
    *,
    page_from: Optional[int] = None,
    page_to: Optional[int] = None,
    strip_chrome: bool = False,
    strip_header: Optional[bool] = None,
    strip_footer: Optional[bool] = None,
) -> List[Tuple[int, str]]:
    """Keep 1-indexed units in range, optionally stripping repeating chrome."""
    if not items:
        return []
    numbers = [num for num, _ in items]
    start, end = clamp_unit_range(max(numbers), page_from, page_to)
    selected = [(num, text) for num, text in items if start <= num <= end]
    header, footer = resolve_strip_sides(
        strip_chrome=strip_chrome,
        strip_header=strip_header,
        strip_footer=strip_footer,
    )
    if not (header or footer) or not selected:
        return selected
    cleaned = strip_repeating_chrome(
        [text for _, text in selected],
        strip_header=header,
        strip_footer=footer,
    )
    return [(num, text) for (num, _), text in zip(selected, cleaned) if (text or "").strip()]


def extract_docx_body_text(filepath: str) -> str:
    """Word body only — section headers and footers are omitted."""
    import docx

    document = docx.Document(filepath)
    parts: List[str] = []
    for paragraph in document.paragraphs:
        text = (paragraph.text or "").strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [(cell.text or "").strip() for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                parts.append(line)
    return "\n".join(parts)


def _story_paragraphs(part: Any) -> List[str]:
    if part is None:
        return []
    texts: List[str] = []
    for paragraph in getattr(part, "paragraphs", None) or []:
        text = (paragraph.text or "").strip()
        if text:
            texts.append(text)
    return texts


def extract_docx_text(
    filepath: str,
    *,
    strip_header: bool = False,
    strip_footer: bool = False,
) -> str:
    """Word extract with independent header/footer keep-or-skip."""
    if not strip_header and not strip_footer:
        import docx2txt

        return docx2txt.process(filepath) or ""
    body = extract_docx_body_text(filepath)
    if strip_header and strip_footer:
        return body
    import docx

    document = docx.Document(filepath)
    headers: List[str] = []
    footers: List[str] = []
    for section in document.sections:
        if not strip_header:
            headers.extend(_story_paragraphs(section.header))
        if not strip_footer:
            footers.extend(_story_paragraphs(section.footer))
    parts: List[str] = []
    if headers:
        parts.append("\n".join(headers))
    if body:
        parts.append(body)
    if footers:
        parts.append("\n".join(footers))
    return "\n".join(parts)


def extract_pdf_page_text(
    reader: Any,
    *,
    page_from: Optional[int] = None,
    page_to: Optional[int] = None,
    strip_chrome: bool = False,
    strip_header: Optional[bool] = None,
    strip_footer: Optional[bool] = None,
) -> str:
    pages = list(getattr(reader, "pages", None) or [])
    total = len(pages)
    if total == 0:
        return ""
    start, end = clamp_unit_range(total, page_from, page_to)
    selected = pages[start - 1 : end]
    texts = [(page.extract_text() or "") for page in selected]
    header, footer = resolve_strip_sides(
        strip_chrome=strip_chrome,
        strip_header=strip_header,
        strip_footer=strip_footer,
    )
    if header or footer:
        texts = strip_repeating_chrome(texts, strip_header=header, strip_footer=footer)
    return "\n".join(texts)
