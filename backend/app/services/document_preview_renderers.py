"""
HTML body renderers for document previews.

Each renderer takes raw file bytes and returns ``(body_html, notice)``. All file
text is HTML-escaped; the only markup comes from this module.
"""
from __future__ import annotations

import base64
import html
import io
import json
from typing import Any, Iterable, List, Optional, Tuple

from app.services.rag.spreadsheet_extract import format_cell

MAX_SHEETS = 20
MAX_SHEET_ROWS = 500
MAX_SHEET_COLS = 50
MAX_COUNTED_ROWS = 100_000
MAX_TEXT_CHARS = 2_000_000
MAX_IMAGE_BYTES = 2 * 1024 * 1024
MAX_TOTAL_IMAGE_BYTES = 12 * 1024 * 1024

RenderResult = Tuple[str, Optional[str]]


def _esc(value: Any) -> str:
    return html.escape(str(value or ""), quote=True)


def _table(rows: Iterable[Iterable[str]]) -> str:
    body = "".join("<tr>" + "".join(f"<td>{_esc(c)}</td>" for c in row) + "</tr>" for row in rows)
    return f'<div class="grid"><table><tbody>{body}</tbody></table></div>' if body else ""


# ── Spreadsheets ──────────────────────────────────────────────────────────────

def _column_label(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index + 1)


def _sheet_table(ws) -> Tuple[str, int, bool]:
    """Return ``(table_html, data_row_count, truncated)`` for one worksheet."""
    ws.reset_dimensions()
    rows: List[Tuple[int, List[str]]] = []
    width = 0
    total = 0
    wide = False
    for row_num, raw in enumerate(ws.iter_rows(values_only=True), start=1):
        values = [format_cell(v) for v in (raw or ())]
        while values and not values[-1]:
            values.pop()
        if not values:
            continue
        total += 1
        if total >= MAX_COUNTED_ROWS:
            break
        if len(rows) < MAX_SHEET_ROWS:
            wide = wide or len(values) > MAX_SHEET_COLS
            values = values[:MAX_SHEET_COLS]
            width = max(width, len(values))
            rows.append((row_num, values))
    if not rows:
        return "", 0, False

    head = '<th class="rownum"></th>' + "".join(f"<th>{_column_label(i)}</th>" for i in range(width))
    body = "".join(
        f'<tr><th class="rownum">{num}</th>'
        + "".join(f"<td>{_esc(v)}</td>" for v in values + [""] * (width - len(values)))
        + "</tr>"
        for num, values in rows
    )
    table = f'<div class="grid"><table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>'
    return table, total, total > len(rows) or wide


def render_spreadsheet(data: bytes) -> RenderResult:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        visible = [ws for ws in wb.worksheets if getattr(ws, "sheet_state", "visible") == "visible"]
        sections: List[str] = []
        links: List[str] = []
        truncated = len(visible) > MAX_SHEETS
        for index, ws in enumerate(visible[:MAX_SHEETS], start=1):
            table, total, cut = _sheet_table(ws)
            truncated = truncated or cut
            anchor = f"sheet-{index}"
            count = f"{total:,}+" if total >= MAX_COUNTED_ROWS else f"{total:,}"
            links.append(f'<a href="#{anchor}">{_esc(ws.title)}</a>')
            content = table or "<p>This sheet is empty.</p>"
            sections.append(
                f'<section class="sheet" id="{anchor}"><h2>{_esc(ws.title)} '
                f'<small>({count} rows)</small></h2>{content}</section>'
            )
    finally:
        wb.close()
    if not sections:
        return "<p>This workbook has no visible sheets.</p>", None
    nav = f'<nav class="sheets">{"".join(links)}</nav>' if len(links) > 1 else ""
    notice = (
        f"Showing the first {MAX_SHEET_ROWS} rows and {MAX_SHEET_COLS} columns of each sheet. "
        "Download the file to see everything."
        if truncated
        else None
    )
    return nav + "".join(sections), notice


# ── Presentations ─────────────────────────────────────────────────────────────

class _ImageBudget:
    def __init__(self) -> None:
        self.used = 0
        self.skipped = 0

    def take(self, blob: bytes) -> bool:
        if len(blob) > MAX_IMAGE_BYTES or self.used + len(blob) > MAX_TOTAL_IMAGE_BYTES:
            self.skipped += 1
            return False
        self.used += len(blob)
        return True


def _shape_html(shape, images: _ImageBudget) -> List[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    parts: List[str] = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            parts.extend(_shape_html(child, images))
        return parts
    if getattr(shape, "has_table", False) and shape.has_table:
        parts.append(_table([(cell.text or "").strip() for cell in row.cells] for row in shape.table.rows))
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            text = "".join(run.text for run in para.runs).strip() or (para.text or "").strip()
            if text:
                indent = f' style="margin-left:{para.level * 1.5:.1f}em"' if para.level else ""
                parts.append(f"<p{indent}>{_esc(text)}</p>")
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        image = shape.image
        if image.content_type.startswith("image/") and images.take(image.blob):
            encoded = base64.b64encode(image.blob).decode("ascii")
            parts.append(f'<img alt="" src="data:{_esc(image.content_type)};base64,{encoded}">')
    return parts


def render_presentation(data: bytes) -> RenderResult:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    images = _ImageBudget()
    sections: List[str] = []
    for number, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title_id = title_shape.shape_id if title_shape is not None else None
        title = (title_shape.text_frame.text or "").strip() if title_shape is not None else ""
        blocks: List[str] = []
        for shape in slide.shapes:
            if shape.shape_id == title_id:
                continue
            try:
                blocks.extend(_shape_html(shape, images))
            except Exception:
                continue
        notes = ""
        if slide.has_notes_slide:
            text = (slide.notes_slide.notes_text_frame.text or "").strip()
            if text:
                notes = f'<div class="notes"><strong>Notes</strong><p>{_esc(text)}</p></div>'
        heading = f"<h2>{_esc(title)}</h2>" if title else ""
        sections.append(
            f'<section class="slide"><div class="slide-num">Slide {number}</div>'
            f"{heading}{''.join(blocks)}{notes}</section>"
        )
    if not sections:
        return "<p>This presentation has no slides.</p>", None
    notice = "Some large images were left out of this preview." if images.skipped else None
    return "".join(sections), notice


# ── Word documents ────────────────────────────────────────────────────────────

def _heading_level(style_name: str) -> Optional[int]:
    name = (style_name or "").lower()
    if name == "title":
        return 1
    if name.startswith("heading "):
        suffix = name.split(" ", 1)[1]
        if suffix.isdigit():
            return min(int(suffix) + 1, 6)
    return None


def render_word(data: bytes) -> RenderResult:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    parts: List[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = (para.text or "").strip()
            if not text:
                continue
            level = _heading_level(para.style.name if para.style is not None else "")
            parts.append(f"<h{level}>{_esc(text)}</h{level}>" if level else f"<p>{_esc(text)}</p>")
        elif tag == "tbl":
            table = Table(child, doc)
            parts.append(_table([(cell.text or "").strip() for cell in row.cells] for row in table.rows))
    if not parts:
        return "<p>This document has no text.</p>", None
    return f'<article class="doc">{"".join(parts)}</article>', None


# ── CSV / TSV ─────────────────────────────────────────────────────────────────

def render_delimited(data: bytes) -> RenderResult:
    import csv

    text = data.decode("utf-8-sig", errors="replace")
    first_line = text.split("\n", 1)[0]
    delimiter = "\t" if first_line.count("\t") > first_line.count(",") else ","
    rows: List[List[str]] = []
    total = 0
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        if not any(cell.strip() for cell in row):
            continue
        total += 1
        if len(rows) < MAX_SHEET_ROWS:
            rows.append(row[:MAX_SHEET_COLS])
    if not rows:
        return "<p>This file is empty.</p>", None
    notice = (
        f"Showing the first {MAX_SHEET_ROWS} of {total:,} rows. Download the file to see everything."
        if total > len(rows)
        else None
    )
    return _table(rows), notice


# ── JSON / XML / other text ───────────────────────────────────────────────────

def render_text(data: bytes) -> RenderResult:
    text = data.decode("utf-8", errors="replace")
    stripped = text.lstrip()
    if stripped.startswith(("{", "[")) and len(text) <= MAX_TEXT_CHARS:
        try:
            text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except ValueError:
            pass
    notice = None
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS]
        notice = "This file is large; only the beginning is shown. Download it to see everything."
    return f"<pre>{_esc(text)}</pre>", notice
