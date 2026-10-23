"""PDF page range and repeating letterhead strip — uploaded documents only."""
from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from pypdf import PdfReader

from app.services.rag.docling_extract import EXTRACT_ENGINE_KEY
from app.services.rag.pdf_page_extract import (
    PDF_PAGE_FROM_KEY,
    PDF_PAGE_TO_KEY,
    PDF_STRIP_CHROME_KEY,
    PDF_STRIP_FOOTER_KEY,
    PDF_STRIP_HEADER_KEY,
    UNSET,
    apply_pdf_options_to_meta,
    extract_pdf_page_text,
    ingest_kwargs_from_document,
    meta_for_upload,
    options_from_meta,
    pdf_fields_for_api,
    resolve_strip_sides,
    strip_repeating_chrome,
    validate_pdf_page_range,
)
from app.services.rag.utils_rag import extract_text_from_file


def _escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _content_stream(lines: list[str]) -> str:
    ops = ["BT", "/F1 12 Tf", "72 720 Td"]
    for index, line in enumerate(lines or [""]):
        if index:
            ops.append("0 -16 Td")
        ops.append(f"({_escape(line)}) Tj")
    ops.append("ET")
    return "\n".join(ops)


def build_text_pdf(pages: list[str]) -> bytes:
    """Minimal multi-page PDF whose text pypdf can extract."""
    font_id = 3
    objs: dict[int, bytes] = {}
    page_ids: list[int] = []
    next_id = 4
    for text in pages:
        page_id = next_id
        content_id = next_id + 1
        next_id += 2
        page_ids.append(page_id)
        stream = _content_stream(text.split("\n"))
        stream_bytes = stream.encode("latin-1")
        objs[content_id] = (
            f"<< /Length {len(stream_bytes)} >>\nstream\n".encode("latin-1")
            + stream_bytes
            + b"\nendstream"
        )
        objs[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {content_id} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>"
        ).encode("latin-1")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objs[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objs[2] = f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode("latin-1")
    objs[font_id] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    max_id = max(objs)
    body = b"%PDF-1.4\n"
    offsets: dict[int, int] = {}
    for i in range(1, max_id + 1):
        offsets[i] = len(body)
        body += f"{i} 0 obj\n".encode("ascii") + objs[i] + b"\nendobj\n"
    xref_pos = len(body)
    xref = [f"xref\n0 {max_id + 1}\n", "0000000000 65535 f \n"]
    for i in range(1, max_id + 1):
        xref.append(f"{offsets[i]:010d} 00000 n \n")
    trailer = (
        f"trailer\n<< /Size {max_id + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n"
    )
    return body + "".join(xref).encode("ascii") + trailer.encode("ascii")


def _letterhead_pages(count: int = 6) -> list[str]:
    pages = []
    for number in range(1, count + 1):
        extra = "Unique Preface Title\n" if number == 1 else ""
        pages.append(
            "ACME Handbook\n"
            f"{extra}"
            f"Section intro {number}\n"
            f"More detail {number}\n"
            f"PAGE_BODY_{number}\n"
            "Confidential\n"
            f"Page {number} of {count}"
        )
    return pages


class _FakePage:
    def __init__(self, text: str):
        self._text = text

    def extract_text(self) -> str:
        return self._text


class _FakeReader:
    def __init__(self, texts: list[str]):
        self.pages = [_FakePage(text) for text in texts]


def test_validate_pdf_page_range_rejects_inverted_and_zero():
    with pytest.raises(ValueError, match="less than or equal"):
        validate_pdf_page_range(5, 2)
    with pytest.raises(ValueError, match=">= 1"):
        validate_pdf_page_range(0, 3)
    assert validate_pdf_page_range(None, None) == (None, None)
    assert validate_pdf_page_range("3", "20") == (3, 20)


def test_options_from_meta_missing_keys_match_current_extract():
    opts = options_from_meta({})
    assert opts.page_from is None and opts.page_to is None
    assert opts.strip_chrome is False
    assert ingest_kwargs_from_document(SimpleNamespace(meta_data=None)) == {}
    fields = pdf_fields_for_api({})
    assert fields == {
        "pdfPageFrom": None,
        "pdfPageTo": None,
        "pdfStripChrome": None,
        "pdfStripHeader": None,
        "pdfStripFooter": None,
        "extractEngine": None,
        "ocrModel": None,
    }
    legacy = options_from_meta({PDF_STRIP_CHROME_KEY: True})
    assert legacy.strip_header is True and legacy.strip_footer is True
    split = options_from_meta({PDF_STRIP_HEADER_KEY: True, PDF_STRIP_FOOTER_KEY: False})
    assert split.strip_header is True and split.strip_footer is False
    assert split.strip_chrome is True
    assert resolve_strip_sides(strip_chrome=True) == (True, True)
    assert resolve_strip_sides(strip_chrome=True, strip_header=True, strip_footer=False) == (
        True,
        False,
    )
    split_kwargs = ingest_kwargs_from_document(
        SimpleNamespace(meta_data={PDF_STRIP_HEADER_KEY: True, PDF_STRIP_FOOTER_KEY: False})
    )
    assert split_kwargs["pdf_strip_chrome"] is True
    assert split_kwargs["pdf_strip_header"] is True
    assert split_kwargs["pdf_strip_footer"] is False


def test_meta_for_upload_stores_options_per_capable_type():
    assert meta_for_upload(page_from=None, page_to=None, strip_chrome=None, filename="notes.txt") == {}
    docx_meta = meta_for_upload(page_from=3, page_to=20, strip_chrome=None, filename="notes.docx")
    assert PDF_PAGE_FROM_KEY not in docx_meta
    assert docx_meta[PDF_STRIP_CHROME_KEY] is True
    assert docx_meta[PDF_STRIP_HEADER_KEY] is True
    assert docx_meta[PDF_STRIP_FOOTER_KEY] is True
    assert docx_meta[EXTRACT_ENGINE_KEY] == "native"
    xlsx_meta = meta_for_upload(page_from=1, page_to=2, strip_chrome=True, filename="data.xlsx")
    assert xlsx_meta[PDF_PAGE_FROM_KEY] == 1
    assert PDF_STRIP_CHROME_KEY not in xlsx_meta
    assert xlsx_meta[EXTRACT_ENGINE_KEY] == "native"
    meta = meta_for_upload(page_from=3, page_to=20, strip_chrome=None, filename="manual.pdf")
    assert meta[PDF_PAGE_FROM_KEY] == 3
    assert meta[PDF_PAGE_TO_KEY] == 20
    assert meta[PDF_STRIP_CHROME_KEY] is True
    assert meta[PDF_STRIP_HEADER_KEY] is True
    assert meta[PDF_STRIP_FOOTER_KEY] is True
    assert meta[EXTRACT_ENGINE_KEY] == "native"
    header_only = meta_for_upload(
        page_from=None,
        page_to=None,
        strip_chrome=True,
        strip_header=True,
        strip_footer=False,
        filename="manual.pdf",
    )
    assert header_only[PDF_STRIP_HEADER_KEY] is True
    assert header_only[PDF_STRIP_FOOTER_KEY] is False
    assert header_only[PDF_STRIP_CHROME_KEY] is True
    pptx_meta = meta_for_upload(
        page_from=2, page_to=4, strip_chrome=True, filename="deck.pptx", extract_engine="docling"
    )
    assert pptx_meta[PDF_PAGE_FROM_KEY] == 2
    assert pptx_meta[PDF_STRIP_CHROME_KEY] is True
    assert pptx_meta[EXTRACT_ENGINE_KEY] == "docling"


def test_apply_pdf_options_to_meta_detects_change_and_clear():
    meta, changed = apply_pdf_options_to_meta(
        {}, page_from=3, page_to=5, strip_chrome=True
    )
    assert changed is True
    assert meta[PDF_PAGE_FROM_KEY] == 3
    assert meta[PDF_STRIP_HEADER_KEY] is True
    assert meta[PDF_STRIP_FOOTER_KEY] is True
    same, again = apply_pdf_options_to_meta(
        meta, page_from=3, page_to=5, strip_chrome=True
    )
    assert again is False
    cleared, cleared_changed = apply_pdf_options_to_meta(
        same, page_from=None, page_to=None, strip_chrome=UNSET
    )
    assert cleared_changed is True
    assert PDF_PAGE_FROM_KEY not in cleared
    assert PDF_STRIP_CHROME_KEY in cleared
    footer_only, footer_changed = apply_pdf_options_to_meta(
        same, strip_header=False, strip_footer=True
    )
    assert footer_changed is True
    assert footer_only[PDF_STRIP_HEADER_KEY] is False
    assert footer_only[PDF_STRIP_FOOTER_KEY] is True
    assert footer_only[PDF_STRIP_CHROME_KEY] is True


def test_strip_repeating_chrome_keeps_unique_title_drops_letterhead():
    pages = _letterhead_pages(5)
    cleaned = strip_repeating_chrome(pages)
    joined = "\n".join(cleaned)
    assert "ACME Handbook" not in joined
    assert "Confidential" not in joined
    assert "Page 1 of 5" not in joined
    assert "Unique Preface Title" in cleaned[0]
    assert "PAGE_BODY_3" in cleaned[2]


def test_strip_repeating_chrome_header_and_footer_are_independent():
    pages = _letterhead_pages(5)
    header_only = "\n".join(strip_repeating_chrome(pages, strip_header=True, strip_footer=False))
    assert "ACME Handbook" not in header_only
    assert "Confidential" in header_only
    assert "PAGE_BODY_3" in header_only

    footer_only = "\n".join(strip_repeating_chrome(pages, strip_header=False, strip_footer=True))
    assert "ACME Handbook" in footer_only
    assert "Confidential" not in footer_only
    assert "Page 1 of 5" not in footer_only
    assert "PAGE_BODY_2" in footer_only

    neither = "\n".join(strip_repeating_chrome(pages, strip_header=False, strip_footer=False))
    assert "ACME Handbook" in neither and "Confidential" in neither


def test_extract_pdf_page_text_range_and_optional_strip():
    pages = _letterhead_pages(8)
    reader = _FakeReader(pages)
    ranged = extract_pdf_page_text(reader, page_from=3, page_to=5, strip_chrome=False)
    assert "PAGE_BODY_3" in ranged and "PAGE_BODY_5" in ranged
    assert "PAGE_BODY_2" not in ranged and "PAGE_BODY_6" not in ranged
    assert "ACME Handbook" in ranged

    stripped = extract_pdf_page_text(reader, page_from=3, page_to=5, strip_chrome=True)
    assert "PAGE_BODY_3" in stripped and "PAGE_BODY_4" in stripped
    assert "ACME Handbook" not in stripped
    assert "Confidential" not in stripped

    header_only = extract_pdf_page_text(
        reader, page_from=3, page_to=5, strip_header=True, strip_footer=False
    )
    assert "ACME Handbook" not in header_only
    assert "Confidential" in header_only


def test_extract_text_from_file_defaults_keep_all_pages_and_letterhead(tmp_path: Path):
    path = tmp_path / "handbook.pdf"
    path.write_bytes(build_text_pdf(_letterhead_pages(6)))
    reader = PdfReader(str(path))
    raw = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "PAGE_BODY_1" in raw and "PAGE_BODY_6" in raw

    texts, _metas = extract_text_from_file(str(path))
    joined = "\n".join(texts)
    assert "PAGE_BODY_1" in joined and "PAGE_BODY_6" in joined
    assert "ACME Handbook" in joined

    ranged, _ = extract_text_from_file(str(path), page_from=3, page_to=5)
    ranged_text = "\n".join(ranged)
    assert "PAGE_BODY_3" in ranged_text and "PAGE_BODY_5" in ranged_text
    assert "PAGE_BODY_1" not in ranged_text and "PAGE_BODY_6" not in ranged_text

    stripped, _ = extract_text_from_file(
        str(path), page_from=1, page_to=6, strip_chrome=True
    )
    stripped_text = "\n".join(stripped)
    assert "PAGE_BODY_2" in stripped_text
    assert "Unique Preface Title" in stripped_text
    assert "ACME Handbook" not in stripped_text
    assert "Confidential" not in stripped_text


def _make_pptx(path: Path, slides: list[str]) -> str:
    from pptx import Presentation

    prs = Presentation()
    for text in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        lines = text.split("\n")
        slide.shapes.title.text = lines[0]
        slide.placeholders[1].text = "\n".join(lines[1:])
    prs.save(str(path))
    return str(path)


def test_pptx_slide_range_and_repeating_chrome(tmp_path: Path):
    slides = []
    for number in range(1, 6):
        extra = "Unique Opening Slide\n" if number == 1 else ""
        slides.append(f"ACME Deck\n{extra}SLIDE_BODY_{number}\nConfidential")
    path = _make_pptx(tmp_path / "deck.pptx", slides)

    all_texts, metas = extract_text_from_file(path)
    joined = "\n".join(all_texts)
    assert "SLIDE_BODY_1" in joined and "SLIDE_BODY_5" in joined
    assert {m.get("slide_number") for m in metas} == {1, 2, 3, 4, 5}

    ranged, ranged_meta = extract_text_from_file(path, page_from=2, page_to=3)
    ranged_text = "\n".join(ranged)
    assert "SLIDE_BODY_2" in ranged_text and "SLIDE_BODY_3" in ranged_text
    assert "SLIDE_BODY_1" not in ranged_text and "SLIDE_BODY_5" not in ranged_text
    assert {m.get("slide_number") for m in ranged_meta} == {2, 3}

    stripped, _ = extract_text_from_file(path, strip_chrome=True)
    stripped_text = "\n".join(stripped)
    assert "SLIDE_BODY_2" in stripped_text
    assert "Unique Opening Slide" in stripped_text
    assert "ACME Deck" not in stripped_text
    assert "Confidential" not in stripped_text


def test_docx_strip_drops_word_header_keeps_body(tmp_path: Path):
    from docx import Document

    path = tmp_path / "policy.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "ACME LETTERHEAD"
    document.add_paragraph("Refunds are issued within 14 days.")
    document.save(str(path))

    kept, _ = extract_text_from_file(str(path))
    kept_text = "\n".join(kept)
    assert "Refunds are issued within 14 days." in kept_text

    stripped, _ = extract_text_from_file(str(path), strip_chrome=True)
    stripped_text = "\n".join(stripped)
    assert "Refunds are issued within 14 days." in stripped_text
    assert "ACME LETTERHEAD" not in stripped_text


def test_docx_header_and_footer_strip_independently(tmp_path: Path):
    from docx import Document

    path = tmp_path / "policy.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "ACME LETTERHEAD"
    document.sections[0].footer.paragraphs[0].text = "CONFIDENTIAL FOOTER"
    document.add_paragraph("Refunds are issued within 14 days.")
    document.save(str(path))

    header_only, _ = extract_text_from_file(str(path), strip_header=True, strip_footer=False)
    header_text = "\n".join(header_only)
    assert "Refunds are issued within 14 days." in header_text
    assert "ACME LETTERHEAD" not in header_text
    assert "CONFIDENTIAL FOOTER" in header_text

    footer_only, _ = extract_text_from_file(str(path), strip_header=False, strip_footer=True)
    footer_text = "\n".join(footer_only)
    assert "Refunds are issued within 14 days." in footer_text
    assert "ACME LETTERHEAD" in footer_text
    assert "CONFIDENTIAL FOOTER" not in footer_text


def test_xlsx_sheet_range_keeps_default_all_visible_sheets(tmp_path: Path):
    from openpyxl import Workbook

    path = tmp_path / "people.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Contacts"
    ws.append(["Name"])
    ws.append(["Ada"])
    prices = wb.create_sheet("Prices")
    prices.append(["Item"])
    prices.append(["Widget"])
    hidden = wb.create_sheet("Secret")
    hidden.append(["Password"])
    hidden.append(["hunter2"])
    hidden.sheet_state = "hidden"
    wb.save(str(path))

    all_texts, _ = extract_text_from_file(str(path))
    joined = "\n".join(all_texts)
    assert "Sheet: Contacts" in joined and "Sheet: Prices" in joined
    assert "hunter2" not in joined

    second, _ = extract_text_from_file(str(path), page_from=2, page_to=2)
    second_text = "\n".join(second)
    assert "Sheet: Prices" in second_text
    assert "Sheet: Contacts" not in second_text


def _uploaded_pdf(**overrides):
    import uuid
    from datetime import datetime, timezone

    base = dict(
        id=uuid.uuid4(),
        user_id=1,
        title="handbook.pdf",
        description=None,
        type="application/pdf",
        source="manual-uploads",
        language="en",
        status="Indexed",
        chunks=4,
        indexed_at=datetime.now(timezone.utc),
        url="",
        checksum="abc",
        size_kb=12,
        ingest_embedding_target=None,
        use_in_answers=True,
        meta_data={},
    )
    base.update(overrides)
    return SimpleNamespace(**base)


def test_update_document_rejects_inverted_page_range():
    import asyncio
    from unittest.mock import MagicMock

    from fastapi import HTTPException

    from app.platform.module_bootstrap import ensure_ragsuite_modules_path

    ensure_ragsuite_modules_path()
    from ragsuite_modules.documents.backend.routes import (  # noqa: E402
        DocumentMetadataUpdate,
        update_document,
    )

    doc = _uploaded_pdf()
    db = MagicMock()
    db.query.return_value.filter.return_value.first.return_value = doc
    body = DocumentMetadataUpdate(pdf_page_from=9, pdf_page_to=2)

    with pytest.raises(HTTPException) as exc:
        asyncio.run(
            update_document(str(doc.id), body, db=db, current_user=SimpleNamespace(id=1))
        )
    assert exc.value.status_code == 400
    assert "pdf_page_from" in str(exc.value.detail)


def test_update_document_persists_pdf_meta_and_retrains_when_range_changes():
    import asyncio
    from unittest.mock import AsyncMock, MagicMock, patch

    from app.platform.module_bootstrap import ensure_ragsuite_modules_path

    ensure_ragsuite_modules_path()
    from ragsuite_modules.documents.backend.routes import (  # noqa: E402
        DocumentMetadataUpdate,
        update_document,
    )
    from ragsuite_modules.documents.backend.upload_helpers import IngestOutcome  # noqa: E402

    doc = _uploaded_pdf()
    db = MagicMock()
    db.query.return_value.filter.return_value.first.return_value = doc
    body = DocumentMetadataUpdate(pdf_page_from=3, pdf_page_to=20, pdf_strip_chrome=True)
    retrain = AsyncMock(return_value=IngestOutcome("Queued", 0, "queued", True))

    with patch(
        "ragsuite_modules.documents.backend.upload_helpers.retrain_stored_document",
        new=retrain,
    ):
        result = asyncio.run(
            update_document(str(doc.id), body, db=db, current_user=SimpleNamespace(id=1))
        )

    assert doc.meta_data[PDF_PAGE_FROM_KEY] == 3
    assert doc.meta_data[PDF_PAGE_TO_KEY] == 20
    assert doc.meta_data[PDF_STRIP_CHROME_KEY] is True
    assert doc.meta_data[PDF_STRIP_HEADER_KEY] is True
    assert doc.meta_data[PDF_STRIP_FOOTER_KEY] is True
    assert result.pdfPageFrom == 3
    assert result.pdfStripChrome is True
    assert result.pdfStripHeader is True
    assert result.pdfStripFooter is True
    retrain.assert_awaited_once()


def test_update_document_title_only_does_not_retrain():
    import asyncio
    from unittest.mock import AsyncMock, MagicMock, patch

    from app.platform.module_bootstrap import ensure_ragsuite_modules_path

    ensure_ragsuite_modules_path()
    from ragsuite_modules.documents.backend.routes import (  # noqa: E402
        DocumentMetadataUpdate,
        update_document,
    )

    doc = _uploaded_pdf(meta_data={PDF_STRIP_CHROME_KEY: True, PDF_PAGE_FROM_KEY: 3})
    db = MagicMock()
    db.query.return_value.filter.return_value.first.return_value = doc
    retrain = AsyncMock()

    with patch(
        "ragsuite_modules.documents.backend.upload_helpers.retrain_stored_document",
        new=retrain,
    ):
        result = asyncio.run(
            update_document(
                str(doc.id),
                DocumentMetadataUpdate(title="Renamed handbook"),
                db=db,
                current_user=SimpleNamespace(id=1),
            )
        )

    assert result.title == "Renamed handbook"
    assert doc.meta_data[PDF_PAGE_FROM_KEY] == 3
    retrain.assert_not_awaited()


def test_update_document_rejects_invalid_custom_split():
    import asyncio
    from unittest.mock import MagicMock

    from fastapi import HTTPException

    from app.platform.module_bootstrap import ensure_ragsuite_modules_path

    ensure_ragsuite_modules_path()
    from ragsuite_modules.documents.backend.routes import (  # noqa: E402
        DocumentMetadataUpdate,
        update_document,
    )

    doc = _uploaded_pdf()
    db = MagicMock()
    db.query.return_value.filter.return_value.first.return_value = doc
    body = DocumentMetadataUpdate(
        split_mode="custom",
        chunk_size=50,
        chunk_overlap=10,
        chunk_delimiter_preset="paragraph",
    )
    with pytest.raises(HTTPException) as exc:
        asyncio.run(update_document(str(doc.id), body, db=db, current_user=SimpleNamespace(id=1)))
    assert exc.value.status_code == 400


def test_update_document_persists_custom_split_without_retrain():
    import asyncio
    from unittest.mock import AsyncMock, MagicMock, patch

    from app.platform.module_bootstrap import ensure_ragsuite_modules_path

    ensure_ragsuite_modules_path()
    from ragsuite_modules.documents.backend.routes import (  # noqa: E402
        DocumentMetadataUpdate,
        update_document,
    )

    doc = _uploaded_pdf()
    db = MagicMock()
    db.query.return_value.filter.return_value.first.return_value = doc
    retrain = AsyncMock()
    body = DocumentMetadataUpdate(
        split_mode="custom",
        chunk_size=800,
        chunk_overlap=200,
        chunk_delimiter_preset="custom",
        chunk_delimiter_custom="|||",
    )
    with patch(
        "ragsuite_modules.documents.backend.upload_helpers.retrain_stored_document",
        new=retrain,
    ):
        result = asyncio.run(update_document(str(doc.id), body, db=db, current_user=SimpleNamespace(id=1)))
    assert result.splitMode == "custom"
    assert (doc.meta_data or {}).get("split_config", {}).get("delimiter_custom") == "|||"
    retrain.assert_not_awaited()
