"""XLSX extraction: header-labelled rows, sheet-scoped chunks, caps, and reindex suffix."""
from __future__ import annotations

import datetime as dt
import io
import zipfile
from pathlib import Path
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook

from app.services.rag import spreadsheet_extract
from app.services.rag.utils_rag import CHUNK_SIZES, extract_text_from_file
from app.services.reindex_service import reindex_temp_suffix_for_uploaded_doc


def _save(wb: Workbook, path: Path) -> str:
    wb.save(str(path))
    return str(path)


def _contacts_workbook(path: Path) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Contacts"
    ws.append(["Name", "Email", "Joined", "Score"])
    ws.append(["Ada Lovelace", "ada@example.com", dt.datetime(2024, 3, 1), 12.0])
    ws.append([None, None, None, None])
    ws.append(["Alan Turing", "alan@example.com", dt.datetime(2024, 3, 2, 9, 30), 7.25])

    prices = wb.create_sheet("Prices")
    prices.append([101, 4.5, True])
    prices.append([102, 6.0, False])

    hidden = wb.create_sheet("Secret")
    hidden.append(["Password"])
    hidden.append(["hunter2"])
    hidden.sheet_state = "hidden"
    return _save(wb, path)


def test_rows_are_header_labelled_and_sheet_scoped(tmp_path: Path):
    texts, metas = extract_text_from_file(_contacts_workbook(tmp_path / "people.xlsx"))

    assert len(texts) == 2
    contacts, prices = texts
    assert contacts.startswith("Sheet: Contacts\n")
    assert "Row 2: Name: Ada Lovelace; Email: ada@example.com; Joined: 2024-03-01; Score: 12" in contacts
    assert "Row 4: Name: Alan Turing; Email: alan@example.com; Joined: 2024-03-02 09:30; Score: 7.25" in contacts
    assert "Row 3" not in contacts

    assert prices.startswith("Sheet: Prices\n")
    assert "Row 1: A: 101; B: 4.5; C: TRUE" in prices
    assert "Row 2: A: 102; B: 6; C: FALSE" in prices

    assert "hunter2" not in "\n".join(texts)
    assert [m["sheet_name"] for m in metas] == ["Contacts", "Prices"]
    assert metas[0]["row_start"] == 2 and metas[0]["row_end"] == 4
    assert metas[1]["sheet_index"] == 2 and metas[1]["page"] == 2
    assert [m["chunk_index"] for m in metas] == [0, 1]
    assert all(m["url"] == "file://people.xlsx" and m["truncated"] is False for m in metas)


def test_rows_pack_until_chunk_size_without_splitting(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Id", "Note"])
    for i in range(60):
        ws.append([i, "x" * 80])
    texts, metas = extract_text_from_file(_save(wb, tmp_path / "notes.xlsx"))

    assert len(texts) > 1
    assert all(len(t) <= CHUNK_SIZES[".xlsx"] for t in texts)
    assert all(t.startswith("Sheet: Sheet\n") for t in texts)
    covered = [r for m in metas for r in range(m["row_start"], m["row_end"] + 1)]
    assert covered == list(range(2, 62))


def test_oversized_row_is_split_with_prefix(tmp_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Title", "Body"])
    long_body = " ".join(f"Sentence number {i} explains the refund policy." for i in range(120))
    ws.append(["Policy", long_body])
    texts, metas = extract_text_from_file(_save(wb, tmp_path / "policy.xlsx"))

    assert len(texts) > 1
    assert all(t.startswith("Sheet: Sheet\nRow 2 (part ") for t in texts)
    assert all(m["row_start"] == m["row_end"] == 2 for m in metas)


def test_row_cap_truncates_and_flags_every_chunk(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(spreadsheet_extract, "MAX_XLSX_DATA_ROWS", 3)
    wb = Workbook()
    ws = wb.active
    ws.append(["N"])
    for i in range(10):
        ws.append([i])
    extra = wb.create_sheet("More")
    extra.append(["ignored"])
    texts, metas = extract_text_from_file(_save(wb, tmp_path / "big.xlsx"))

    joined = "\n".join(texts)
    assert "Row 4: N: 2" in joined
    assert "Row 5" not in joined
    assert "More" not in joined
    assert metas and all(m["truncated"] is True for m in metas)


def test_exact_cap_is_not_marked_truncated(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(spreadsheet_extract, "MAX_XLSX_DATA_ROWS", 2)
    wb = Workbook()
    ws = wb.active
    ws.append(["N"])
    ws.append([1])
    ws.append([2])
    _texts, metas = extract_text_from_file(_save(wb, tmp_path / "exact.xlsx"))
    assert all(m["truncated"] is False for m in metas)


def test_empty_workbook_yields_no_chunks(tmp_path: Path):
    texts, metas = extract_text_from_file(_save(Workbook(), tmp_path / "empty.xlsx"))
    assert texts == [] and metas == []


def test_header_only_sheet_yields_no_chunks(tmp_path: Path):
    wb = Workbook()
    wb.active.append(["Name", "Email"])
    texts, _ = extract_text_from_file(_save(wb, tmp_path / "header.xlsx"))
    assert texts == []


def test_corrupt_xlsx_raises_value_error(tmp_path: Path):
    bad = tmp_path / "broken.xlsx"
    bad.write_bytes(b"not a zip at all")
    with pytest.raises(ValueError, match="Not a valid XLSX workbook"):
        extract_text_from_file(str(bad))


def test_legacy_xls_is_rejected(tmp_path: Path):
    legacy = tmp_path / "old.xls"
    legacy.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(ValueError, match=r"\.xls"):
        extract_text_from_file(str(legacy))


def test_format_cell_handles_common_types():
    fmt = spreadsheet_extract.format_cell
    assert fmt(None) == ""
    assert fmt(3.0) == "3"
    assert fmt(0.1) == "0.1"
    assert fmt(dt.date(2024, 1, 5)) == "2024-01-05"
    assert fmt("  two\nlines\t here ") == "two lines here"
    assert fmt("bad\ud800char") == "badchar"


def test_reindex_suffix_maps_spreadsheet_zip_to_xlsx():
    doc = MagicMock()
    doc.title = "Quarterly numbers"
    doc.type = "application/XLSX"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("xl/workbook.xml", "<workbook/>")
    assert reindex_temp_suffix_for_uploaded_doc(doc, buf.getvalue()) == ".xlsx"


def test_reindex_suffix_uses_xlsx_title():
    doc = MagicMock()
    doc.title = "prices.xlsx"
    doc.type = "application/XLSX"
    assert reindex_temp_suffix_for_uploaded_doc(doc, b"PK\x03\x04") == ".xlsx"
