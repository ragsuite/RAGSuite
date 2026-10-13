"""Structure capture (headings, breadcrumbs) and contextual embedding at ingest."""
from __future__ import annotations

from pathlib import Path

from bs4 import BeautifulSoup

from app.services.html_page_text import extract_page_text_with_outline, json_ld_breadcrumb
from app.services.rag import citation_trailer, singleton
from app.services.rag.utils_rag import extract_text_from_file

PAGE = """
<html><head>
<script type="application/ld+json">
{"@context": "https://schema.org", "@graph": [{"@type": "BreadcrumbList", "itemListElement": [
  {"@type": "ListItem", "position": 2, "name": "Hand cleaning"},
  {"@type": "ListItem", "position": 1, "item": {"name": "Products"}},
  {"@type": "ListItem", "position": 3, "name": "Soap dispenser"}]}]}
</script></head>
<body><main>
<h1>Soap dispenser SD-500</h1><p>Wall mounted dispenser for liquid soap.</p>
<h2>Technical data</h2><p>Capacity 1000 ml, stainless steel.</p>
</main></body></html>
"""


def test_html_outline_and_json_ld_breadcrumb():
    soup = BeautifulSoup(PAGE, "html.parser")
    assert json_ld_breadcrumb(soup) == "Products > Hand cleaning > Soap dispenser"
    text, headings = extract_page_text_with_outline(soup.find("main"))
    assert headings == ["Soap dispenser SD-500", "Technical data"]
    assert "Capacity 1000 ml" in text


def test_json_ld_breadcrumb_absent_or_broken():
    assert json_ld_breadcrumb(BeautifulSoup("<html><body>x</body></html>", "html.parser")) == ""
    broken = '<script type="application/ld+json">{not json</script>'
    assert json_ld_breadcrumb(BeautifulSoup(broken, "html.parser")) == ""


def test_label_passages_adds_page_line_before_url():
    metas = [{"title": "Soap dispenser", "section": "Technical data"}, {}]
    out = citation_trailer.label_passages(["Capacity 1000 ml", "Other"], metas, lambda m: "Source URL: u\n")
    assert out[0] == "[S1]\nPage: Soap dispenser > Technical data\nSource URL: u\nCapacity 1000 ml"
    assert out[1] == "[S2]\nSource URL: u\nOther"


def test_docx_heading_styles_become_sections(tmp_path: Path):
    import docx

    document = docx.Document()
    document.add_heading("Installation", level=1)
    document.add_paragraph("Mount the bracket on a flat wall. " * 30)
    document.add_heading("Maintenance", level=1)
    document.add_paragraph("Clean the nozzle every week with warm water. " * 30)
    path = tmp_path / "manual.docx"
    document.save(str(path))

    texts, metas = extract_text_from_file(str(path))
    sections = [m.get("section") for m in metas]
    assert sections[0] == "Installation"
    assert sections[-1] == "Maintenance"


def test_pptx_slide_title_becomes_section(tmp_path: Path):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quarterly results"
    slide.placeholders[1].text = "Revenue grew in every region."
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    _, metas = extract_text_from_file(str(path))
    assert metas and all(m.get("section") == "Quarterly results" for m in metas)


class _Embedder:
    def __init__(self):
        self.seen = []

    def embed(self, texts, on_progress=None):
        self.seen.extend(texts)
        return [[0.0] for _ in texts]


class _Pipeline:
    def __init__(self):
        self.embedder = _Embedder()
        self.written = {}

    def _embedder_for(self, provider, model, api_key, endpoint=None, api_version=None):
        return self.embedder

    def write_prepared_ingest(self, **kwargs):
        self.written = kwargs


def test_prepared_ingest_embeds_header_but_stores_raw_text(monkeypatch):
    pipeline = _Pipeline()
    monkeypatch.setattr(singleton, "get_pipeline", lambda: pipeline)
    meta = {"title": "Soap dispenser", "url": "https://x.com/en/products/hand-cleaning/sd.html"}
    result = singleton.locked_write_prepared_ingest(
        texts=["Capacity 1000 ml"], chunk_metadata=[meta], source_file="crawl_source_x", document_id="d1"
    )
    assert result["status"] == "Indexed"
    assert pipeline.embedder.seen == ["Soap dispenser > products > hand cleaning\nCapacity 1000 ml"]
    assert pipeline.written["texts"] == ["Capacity 1000 ml"]
    stored = pipeline.written["chunk_metadata"][0]
    assert stored["chunk_ctx_v"] == 1 and stored["ctx_header"].startswith("Soap dispenser")
    assert "ctx_header" not in meta
