"""Browser previews for documents browsers would otherwise download (XLSX / PPTX / DOCX / JSON)."""
import io
import uuid
from datetime import timedelta
from types import SimpleNamespace

import pytest

from app.platform.module_bootstrap import ensure_ragsuite_modules_path
from app.services.document_preview_html import (
    KIND_DELIMITED,
    KIND_HTML,
    KIND_NATIVE,
    KIND_PRESENTATION,
    KIND_SPREADSHEET,
    KIND_TEXT,
    KIND_UNSUPPORTED,
    KIND_WORD,
    preview_kind,
    render_preview_html,
)
from app.services import document_preview_renderers as renderers

ensure_ragsuite_modules_path()

from ragsuite_modules.documents.backend.routes import (  # noqa: E402
    _document_media_type,
    get_document_content_stream,
)

XSS = "<script>alert(1)</script>"


def _xlsx_bytes(rows, *, extra_sheet=False) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "People"
    for row in rows:
        ws.append(row)
    if extra_sheet:
        wb.create_sheet("Empty")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx_bytes() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quarterly review"
    slide.placeholders[1].text_frame.text = f"Revenue grew {XSS}"
    slide.notes_slide.notes_text_frame.text = "Speaker notes here"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Handbook", level=1)
    doc.add_paragraph(f"Welcome {XSS}")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Key"
    table.rows[0].cells[1].text = "Value"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize(
    "media_type,filename,expected",
    [
        ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "a.xlsx", KIND_SPREADSHEET),
        ("application/xlsx", "1mb.xlsx", KIND_SPREADSHEET),
        ("application/octet-stream", "deck.pptx", KIND_PRESENTATION),
        ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "x", KIND_WORD),
        ("application/json", "data", KIND_TEXT),
        ("application/x-ragsuite-qa+json", "QA", KIND_TEXT),
        ("text/html", "page.html", KIND_HTML),
        ("application/pdf", "a.pdf", KIND_NATIVE),
        ("text/plain", "a.txt", KIND_NATIVE),
        ("text/csv", "a.csv", KIND_DELIMITED),
        ("application/vnd.ms-excel", "export.csv", KIND_DELIMITED),
        ("text/markdown", "notes.md", KIND_TEXT),
        ("application/msword", "old.doc", KIND_UNSUPPORTED),
        ("application/vnd.ms-excel", "old.xls", KIND_UNSUPPORTED),
    ],
)
def test_preview_kind(media_type, filename, expected):
    assert preview_kind(media_type, filename) == expected


def test_media_type_expands_truncated_office_types():
    assert _document_media_type("application/XLSX").endswith("spreadsheetml.sheet")
    assert _document_media_type("application/PPTX").endswith("presentationml.presentation")
    assert _document_media_type("application/DOCX").endswith("wordprocessingml.document")
    assert _document_media_type("application/x-ragsuite-qa+json") == "application/x-ragsuite-qa+json"


def test_spreadsheet_preview_renders_grid_and_escapes_cells():
    body, notice = renderers.render_spreadsheet(
        _xlsx_bytes([["Name", "Role"], ["Ada", XSS], [None, None], ["Linus", 42]], extra_sheet=True)
    )
    assert notice is None
    assert "<th>A</th>" in body and "<th>B</th>" in body
    assert "<td>Ada</td>" in body and "<td>42</td>" in body
    assert XSS not in body and "&lt;script&gt;" in body
    assert 'href="#sheet-2"' in body and "This sheet is empty." in body
    assert "(3 rows)" in body


def test_spreadsheet_preview_truncates_long_sheets(monkeypatch):
    monkeypatch.setattr(renderers, "MAX_SHEET_ROWS", 2)
    body, notice = renderers.render_spreadsheet(_xlsx_bytes([[i] for i in range(5)]))
    assert body.count('<th class="rownum">') == 3  # header corner + 2 rows
    assert "(5 rows)" in body
    assert notice and "first 2 rows" in notice


def test_presentation_preview_lists_slides_and_notes():
    body, _ = renderers.render_presentation(_pptx_bytes())
    assert "Slide 1" in body and "<h2>Quarterly review</h2>" in body
    assert "Revenue grew &lt;script&gt;" in body and XSS not in body
    assert "Speaker notes here" in body


def test_word_preview_keeps_headings_and_tables():
    body, _ = renderers.render_word(_docx_bytes())
    assert "<h2>Handbook</h2>" in body
    assert "<td>Key</td><td>Value</td>" in body
    assert XSS not in body


def test_delimited_preview_renders_csv_and_tsv_tables():
    body, notice = renderers.render_delimited(b'name,role\nAda,"CTO, <b>"\n\n')
    assert "<td>Ada</td><td>CTO, &lt;b&gt;</td>" in body and notice is None
    tsv, _ = renderers.render_delimited(b"a\tb\n1\t2\n")
    assert "<td>1</td><td>2</td>" in tsv


def test_text_preview_pretty_prints_json():
    body, _ = renderers.render_text(b'{"a":[1,2],"b":"<i>x</i>"}')
    assert "&quot;a&quot;: [" in body and "&lt;i&gt;" in body


def test_page_falls_back_to_download_for_unsupported_and_broken_files():
    page = render_preview_html(b"x", title=f"old {XSS}.doc", kind=KIND_UNSUPPORTED, download_url="content-stream?token=t&download=1")
    assert "preview isn&#x27;t available" in page.lower()
    assert 'href="content-stream?token=t&amp;download=1" download' in page
    assert XSS not in page
    broken = render_preview_html(b"not a zip", title="x.xlsx", kind=KIND_SPREADSHEET, download_url="d")
    assert "couldn&#x27;t be previewed" in broken
    assert "<script" not in broken


class _Query:
    def __init__(self, row):
        self.row = row

    def filter(self, *args, **kwargs):
        return self

    def first(self):
        return self.row


def _stream(doc, *, fetch_dest=None, embed=False, **params):
    from app.auth import create_access_token

    token = create_access_token(
        {"doc_id": str(doc.id), "scope": "content", "sub": "tester"}, timedelta(minutes=5)
    )
    db = SimpleNamespace(query=lambda *_: _Query(doc), commit=lambda: None)
    request = SimpleNamespace(headers={"sec-fetch-dest": fetch_dest} if fetch_dest else {})
    return get_document_content_stream(
        str(doc.id), request=request, token=token, db=db, embed=embed, **params
    )


def _doc(content: bytes, *, type_: str, title: str):
    return SimpleNamespace(id=uuid.uuid4(), text_content=content, type=type_, title=title, user_id=1)


def test_content_stream_preview_download_and_default():
    doc = _doc(_xlsx_bytes([["Name"], ["Ada"]]), type_="application/XLSX", title="1mb.xlsx")

    preview = _stream(doc, preview=True, download=False)
    assert preview.media_type == "text/html"
    assert b"<td>Ada</td>" in preview.body
    assert b"<header>" in preview.body
    embedded = _stream(doc, preview=True, download=False, embed=True)
    assert b"<header>" not in embedded.body and b"<td>Ada</td>" in embedded.body
    assert "default-src 'none'" in preview.headers["content-security-policy"]
    assert preview.headers["cache-control"] == "no-store"

    download = _stream(doc, preview=False, download=True)
    assert download.headers["content-disposition"].startswith("attachment;")
    assert download.media_type.endswith("spreadsheetml.sheet")

    default = _stream(doc, preview=False, download=False)
    assert default.headers["content-disposition"].startswith("attachment;")


def test_content_stream_browser_navigation_gets_preview_unless_downloading():
    """Widget citations open content-stream without ``preview=1``; they must not download."""
    doc = _doc(_xlsx_bytes([["Name"], ["Ada"]]), type_="application/XLSX", title="1mb.xlsx")

    for dest in ("document", "iframe"):
        page = _stream(doc, fetch_dest=dest, preview=False, download=False)
        assert page.media_type == "text/html" and b"<td>Ada</td>" in page.body

    fetched = _stream(doc, fetch_dest="empty", preview=False, download=False)
    assert fetched.headers["content-disposition"].startswith("attachment;")

    download = _stream(doc, fetch_dest="document", preview=False, download=True)
    assert download.headers["content-disposition"].startswith("attachment;")


def test_content_stream_preview_keeps_pdf_native_and_sandboxes_html():
    pdf = _stream(_doc(b"%PDF-1.4", type_="application/pdf", title="a.pdf"), preview=True, download=False)
    assert pdf.headers["content-disposition"].startswith("inline;")

    page = _stream(_doc(b"<b>hi</b>", type_="text/html", title="a.html"), preview=True, download=False)
    assert page.headers["content-security-policy"].startswith("sandbox")
